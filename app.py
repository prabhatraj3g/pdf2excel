from __future__ import annotations

import io
import json
import re
import sqlite3
from copy import copy
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
import streamlit as st
from PIL import Image, ImageOps
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image as PdfImage,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


SUPPORTED_TYPES = {
    "pdf": ["pdf"],
    "image": ["png", "jpg", "jpeg", "webp", "bmp", "tiff"],
    "spreadsheet": ["xlsx", "xls", "csv"],
    "document": ["docx", "txt", "md"],
}
DATABASE_PATH = Path(__file__).with_name("praj_pdf_converter.db")


@dataclass
class ConvertedFile:
    name: str
    kind: str
    pages: int = 1
    ocr_text: str = ""


def initialize_database() -> None:
    with sqlite3.connect(DATABASE_PATH) as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS conversion_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                status TEXT NOT NULL,
                output_name TEXT NOT NULL,
                file_count INTEGER NOT NULL,
                ocr_enabled INTEGER NOT NULL,
                ocr_completed INTEGER NOT NULL,
                files_json TEXT NOT NULL,
                error_message TEXT
            )
            """
        )


def log_conversion(
    status: str,
    output_name: str,
    files,
    include_ocr: bool,
    converted: list[ConvertedFile] | None = None,
    error_message: str = "",
) -> None:
    converted_by_name = {item.name: item for item in converted or []}
    file_details = []
    for uploaded in files:
        converted_item = converted_by_name.get(uploaded.name)
        file_details.append(
            {
                "name": uploaded.name,
                "kind": file_kind(uploaded.name) or "unsupported",
                "size_bytes": uploaded.size,
                "pages": converted_item.pages if converted_item else 0,
                "ocr_completed": bool(converted_item and converted_item.ocr_text and "could not run" not in converted_item.ocr_text),
            }
        )
    ocr_completed = sum(item["ocr_completed"] for item in file_details)
    with sqlite3.connect(DATABASE_PATH) as connection:
        connection.execute(
            """
            INSERT INTO conversion_logs
            (created_at, status, output_name, file_count, ocr_enabled, ocr_completed, files_json, error_message)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                datetime.now(timezone.utc).isoformat(timespec="seconds"),
                status,
                output_name,
                len(files),
                int(include_ocr),
                ocr_completed,
                json.dumps(file_details),
                error_message[:1000],
            ),
        )


def recent_logs(limit: int = 8) -> list[tuple]:
    with sqlite3.connect(DATABASE_PATH) as connection:
        return connection.execute(
            """
            SELECT created_at, status, output_name, file_count, ocr_completed, error_message
            FROM conversion_logs
            ORDER BY id DESC
            LIMIT ?
            """,
            (limit,),
        ).fetchall()


def html_text(value: object) -> str:
    text = "" if value is None else str(value)
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("\n", "<br/>")
    )


def file_kind(name: str) -> str | None:
    suffix = Path(name).suffix.lower().lstrip(".")
    for kind, extensions in SUPPORTED_TYPES.items():
        if suffix in extensions:
            return kind
    return None


def safe_ocr(image: Image.Image) -> str:
    try:
        import pytesseract

        prepared = ImageOps.exif_transpose(image).convert("L")
        prepared = ImageOps.autocontrast(prepared)
        if max(prepared.size) < 1800:
            scale = 1800 / max(prepared.size)
            prepared = prepared.resize((int(prepared.width * scale), int(prepared.height * scale)))
        return pytesseract.image_to_string(prepared, config="--psm 3").strip()
    except Exception:
        return "OCR could not run because the Tesseract engine is not installed."


def add_header_footer(canvas, document) -> None:
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#d8dee9"))
    canvas.line(18 * mm, 15 * mm, 192 * mm, 15 * mm)
    canvas.setFillColor(colors.HexColor("#667085"))
    canvas.setFont("Helvetica", 8)
    canvas.drawString(18 * mm, 10 * mm, "Praj PDF Converter workspace")
    canvas.drawRightString(192 * mm, 10 * mm, f"Page {document.page}")
    canvas.restoreState()


def table_from_dataframe(dataframe: pd.DataFrame) -> Table:
    frame = dataframe.fillna("").astype(str).iloc[:100, :12]
    rows = [[html_text(column) for column in frame.columns]] + [
        [html_text(value) for value in row] for row in frame.itertuples(index=False, name=None)
    ]
    table = Table(rows, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17324d")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f7fa")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def build_pdf(files, include_ocr: bool) -> tuple[bytes, list[ConvertedFile]]:
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CoverTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=25, leading=30, textColor=colors.HexColor("#17324d"), alignment=TA_CENTER, spaceAfter=12))
    styles.add(ParagraphStyle(name="SectionTitle", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=colors.HexColor("#17324d"), spaceBefore=8, spaceAfter=8))
    styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontSize=8, leading=11, textColor=colors.HexColor("#475467")))
    styles.add(ParagraphStyle(name="BodyTight", parent=styles["BodyText"], fontSize=9, leading=13, spaceAfter=5))

    story = []
    converted: list[ConvertedFile] = []

    for index, uploaded in enumerate(files, start=1):
        name = uploaded.name
        kind = file_kind(name)
        if kind is None:
            continue
        converted_item = ConvertedFile(name=name, kind=kind)
        if kind != "image":
            story.append(Paragraph(f"{index:02d}  {html_text(name)}", styles["SectionTitle"]))
            story.append(Paragraph(f"Source type: {kind.title()}", styles["Small"]))
            story.append(Spacer(1, 4 * mm))
        if kind == "image":
            image = Image.open(uploaded).convert("RGB")
            image.thumbnail((165 * mm, 210 * mm))
            image_buffer = io.BytesIO()
            image.save(image_buffer, format="PNG")
            image_buffer.seek(0)
            story.append(PdfImage(image_buffer, width=image.width * 0.45, height=image.height * 0.45))
            if include_ocr:
                converted_item.ocr_text = safe_ocr(image)
                if converted_item.ocr_text and "could not run" not in converted_item.ocr_text:
                    story.extend([Spacer(1, 4 * mm), Paragraph("OCR text", styles["Heading3"]), Paragraph(html_text(converted_item.ocr_text), styles["BodyTight"])])
        elif kind == "spreadsheet":
            if Path(name).suffix.lower() == ".csv":
                dataframe = pd.read_csv(uploaded)
                story.append(table_from_dataframe(dataframe))
                story.append(Paragraph(f"Showing up to 100 rows and 12 columns from {len(dataframe):,} row(s).", styles["Small"]))
            else:
                uploaded.seek(0)
                workbook = pd.ExcelFile(uploaded)
                for sheet_number, sheet_name in enumerate(workbook.sheet_names):
                    if sheet_number:
                        story.append(PageBreak())
                    dataframe = pd.read_excel(workbook, sheet_name=sheet_name)
                    story.append(Paragraph(f"Sheet: {html_text(sheet_name)}", styles["Heading3"]))
                    story.append(table_from_dataframe(dataframe))
                    story.append(Paragraph(f"Showing up to 100 rows and 12 columns from {len(dataframe):,} row(s).", styles["Small"]))
        elif kind == "document":
            if Path(name).suffix.lower() == ".docx":
                from docx import Document

                document = Document(uploaded)
                paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
                for paragraph in paragraphs:
                    story.append(Paragraph(html_text(paragraph), styles["BodyTight"]))
            else:
                text = uploaded.getvalue().decode("utf-8", errors="replace")
                for paragraph in re.split(r"\n\s*\n", text):
                    if paragraph.strip():
                        story.append(Paragraph(html_text(paragraph), styles["BodyTight"]))
        elif kind == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(uploaded)
            converted_item.pages = len(reader.pages)
            story.append(Paragraph(f"Existing PDF with {converted_item.pages} page(s). It will be included in the final package.", styles["BodyTight"]))
            story.append(Paragraph("Note: PDF pages are listed here; direct page merging is supported when pypdf is available in the runtime.", styles["Small"]))

        converted.append(converted_item)
        if index != len(files):
            story.append(PageBreak())

    SimpleDocTemplate(output, pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=18 * mm, bottomMargin=20 * mm, title="Praj PDF Converter").build(story, onFirstPage=add_header_footer, onLaterPages=add_header_footer)

    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for page in PdfReader(io.BytesIO(output.getvalue())).pages:
        writer.add_page(page)
    for uploaded in files:
        if file_kind(uploaded.name) == "pdf":
            uploaded.seek(0)
            for page in PdfReader(uploaded).pages:
                writer.add_page(page)
    merged_output = io.BytesIO()
    writer.write(merged_output)
    return merged_output.getvalue(), converted


def excel_sheet_name(value: str, used_names: set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", value).strip() or "Table"
    cleaned = cleaned[:31]
    candidate = cleaned
    suffix = 2
    while candidate in used_names:
        suffix_text = f"_{suffix}"
        candidate = f"{cleaned[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    used_names.add(candidate)
    return candidate


def image_table(image: Image.Image) -> pd.DataFrame:
    import pytesseract
    from pytesseract import Output

    prepared = ImageOps.exif_transpose(image).convert("RGB")
    prepared = ImageOps.autocontrast(prepared.convert("L"))
    if max(prepared.size) < 1800:
        scale = 1800 / max(prepared.size)
        prepared = prepared.resize((int(prepared.width * scale), int(prepared.height * scale)))
    words = pytesseract.image_to_data(prepared, config="--psm 6", output_type=Output.DICT)
    tokens = []
    for index, text in enumerate(words["text"]):
        text = text.strip()
        if text and float(words["conf"][index]) >= 20:
            tokens.append((int(words["left"][index]), int(words["top"][index]), text))
    if not tokens:
        return pd.DataFrame()

    rows: list[list[tuple[int, str]]] = []
    row_tolerance = max(12, prepared.height // 100)
    for x, y, text in sorted(tokens, key=lambda item: (item[1], item[0])):
        row = next((candidate for candidate in rows if abs(candidate[0][1] - y) <= row_tolerance), None)
        if row is None:
            rows.append([(y, ""), (x, text)])
        else:
            row.append((x, text))
    rows.sort(key=lambda row: row[0][0])
    values = [[text for _, text in sorted(row[1:], key=lambda item: item[0])] for row in rows]
    width = max(len(row) for row in values)
    return pd.DataFrame([row + [""] * (width - len(row)) for row in values])


def bank_statement_tables(uploaded, filename: str) -> list[tuple[str, pd.DataFrame]]:
    import pdfplumber

    columns = ["Trans Date", "Value Date", "Branch", "Ref/Chq No", "Description", "Withdrawals (Dr)", "Deposit (Cr)", "Balance"]
    date_pattern = re.compile(r"^\d{1,2}[/-]\d{1,2}[/-]\d{2,4}$")
    amount_pattern = re.compile(r"^[\d,]+(?:\.\d{1,2})?$")
    tables: list[tuple[str, pd.DataFrame]] = []
    uploaded.seek(0)
    with pdfplumber.open(uploaded) as document:
        for page_number, page in enumerate(document.pages, start=1):
            words = page.extract_words(x_tolerance=2, y_tolerance=3)
            if not words:
                continue
            lines: list[list[dict]] = []
            for word in sorted(words, key=lambda item: (item["top"], item["x0"])):
                line = next((candidate for candidate in lines if abs(candidate[0]["top"] - word["top"]) <= 3), None)
                if line is None:
                    lines.append([word])
                else:
                    line.append(word)
            lines = [sorted(line, key=lambda item: item["x0"]) for line in lines]
            header = next((line for line in lines if sum(word["text"].lower().startswith(term) for word in line for term in ["trans", "value", "branch", "description", "balance"]) >= 3), None)
            if header is None:
                continue
            header_text = " ".join(word["text"].lower() for word in header)
            if "withdraw" not in header_text and "deposit" not in header_text:
                continue
            starts = []
            labels = [("trans", "Trans Date"), ("value", "Value Date"), ("branch", "Branch"), ("ref", "Ref/Chq No"), ("description", "Description"), ("withdraw", "Withdrawals (Dr)"), ("deposit", "Deposit (Cr)"), ("balance", "Balance")]
            for term, _ in labels:
                match = next((word for word in header if word["text"].lower().startswith(term)), None)
                if match:
                    starts.append((match["x0"], term))
            starts.sort()
            if len(starts) < 6:
                continue
            boundaries = [position for position, _ in starts]
            rows: list[dict[str, str]] = []
            for line in lines[lines.index(header) + 1:]:
                cells = {term: [] for _, term in starts}
                for word in line:
                    index = min(range(len(boundaries)), key=lambda item: abs(word["x0"] - boundaries[item]))
                    cells[starts[index][1]].append(word["text"])
                values = {term: " ".join(items).strip() for term, items in cells.items()}
                if not any(values.values()):
                    continue
                first_value = next(iter(values.values()), "")
                if not date_pattern.match(first_value):
                    if rows and values.get("description"):
                        rows[-1]["Description"] = f"{rows[-1]['Description']} {values['description']}".strip()
                    continue
                row = {column: "" for column in columns}
                row["Trans Date"] = first_value
                for term, column in labels[1:]:
                    if term in values:
                        row[column] = values[term]
                rows.append(row)
            if rows:
                tables.append((f"{Path(filename).stem} Page {page_number}", pd.DataFrame(rows, columns=columns)))
    return tables


def pdf_text_tables(uploaded, filename: str) -> list[tuple[str, pd.DataFrame]]:
    import pdfplumber

    page_rows: list[dict[str, object]] = []
    field_rows_by_table: dict[str, list[dict[str, str]]] = {}
    current_table = ""
    current_field: dict[str, str] | None = None
    column_starts: list[float] | None = None

    def grouped_lines(words: list[dict]) -> list[list[dict]]:
        lines: list[list[dict]] = []
        for word in sorted(words, key=lambda item: (item["top"], item["x0"])):
            line = next((candidate for candidate in lines if abs(candidate[0]["top"] - word["top"]) <= 3), None)
            if line is None:
                lines.append([word])
            else:
                line.append(word)
        return [sorted(line, key=lambda item: item["x0"]) for line in lines]

    def column_values(line: list[dict]) -> list[str]:
        if not column_starts:
            return []
        values = [[] for _ in column_starts]
        for word in line:
            column_index = min(range(len(column_starts)), key=lambda index: abs(word["x0"] - column_starts[index]))
            values[column_index].append(word["text"])
        return [" ".join(value).strip() for value in values]

    uploaded.seek(0)
    with pdfplumber.open(uploaded) as document:
        for page_number, page in enumerate(document.pages, start=1):
            words = page.extract_words(x_tolerance=2, y_tolerance=3)
            lines = grouped_lines(words)
            for line in lines:
                clean_line = " ".join(word["text"] for word in line).strip()
                if not clean_line:
                    continue
                page_rows.append({"Page": page_number, "Text": clean_line})
                table_match = re.match(r"Table:\s*(.+)", clean_line, flags=re.IGNORECASE)
                if table_match:
                    current_table = table_match.group(1).strip()
                    current_field = None
                    field_rows_by_table.setdefault(current_table, [])
                    continue
                header_words = [word["text"].lower() for word in line]
                if "name" in header_words and "size" in header_words and "description" in header_words:
                    header_positions = {word["text"].lower(): word["x0"] for word in line}
                    type_positions = [word["x0"] for word in line if word["text"].lower() == "type"]
                    column_starts = [
                        header_positions.get("name", 0),
                        type_positions[0] if type_positions else 0,
                        type_positions[1] if len(type_positions) > 1 else 0,
                        header_positions.get("size", 0),
                        header_positions.get("description", 0),
                    ]
                    continue
                values = column_values(line)
                if current_table and column_starts and values and values[0] and values[1] in {"Number", "Text", "Date", "Time", "Boolean"}:
                    current_field = {
                        "Name": values[0],
                        "Data type": values[1],
                        "Format": values[2],
                        "Size": values[3],
                        "Description": values[4],
                    }
                    field_rows_by_table[current_table].append(current_field)
                elif current_field and not clean_line.startswith(("Table Indexes", "KEY_", "Columns", "Name Type")):
                    current_field["Description"] = f"{current_field['Description']} {clean_line}".strip()

    tables = []
    for table_name, field_rows in field_rows_by_table.items():
        if field_rows:
            tables.append((table_name, pd.DataFrame(field_rows, columns=["Table", "Name", "Data type", "Format", "Size", "Description"]).drop(columns=["Table"])))
    tables.append((f"{Path(filename).stem} Text", pd.DataFrame(page_rows, columns=["Page", "Text"])))
    return tables


def build_excel(files) -> tuple[bytes, list[str]]:
    output = io.BytesIO()
    used_sheet_names: set[str] = set()
    sheets_created: list[str] = []
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for uploaded in files:
            name = uploaded.name
            kind = file_kind(name)
            if kind is None:
                continue
            suffix = Path(name).suffix.lower()
            if kind == "spreadsheet":
                uploaded.seek(0)
                if suffix == ".csv":
                    tables = [(Path(name).stem, pd.read_csv(uploaded))]
                else:
                    workbook = pd.ExcelFile(uploaded)
                    tables = [(sheet_name, pd.read_excel(workbook, sheet_name=sheet_name)) for sheet_name in workbook.sheet_names]
            elif kind == "pdf":
                import pdfplumber

                uploaded.seek(0)
                tables = bank_statement_tables(uploaded, name)
                if not tables:
                    uploaded.seek(0)
                    with pdfplumber.open(uploaded) as document:
                        for page_number, page in enumerate(document.pages, start=1):
                            for table_number, table in enumerate(page.extract_tables(), start=1):
                                if table:
                                    frame = pd.DataFrame(table[1:], columns=table[0]) if len(table) > 1 else pd.DataFrame(table)
                                    tables.append((f"Page {page_number} Table {table_number}", frame))
                if not tables:
                    tables = pdf_text_tables(uploaded, name)
            elif kind == "image":
                uploaded.seek(0)
                tables = [(Path(name).stem, image_table(Image.open(uploaded)))]
            elif kind == "document" and suffix == ".docx":
                from docx import Document

                document = Document(uploaded)
                tables = [(f"{Path(name).stem} Table {index}", pd.DataFrame([[cell.text for cell in row.cells] for row in table.rows])) for index, table in enumerate(document.tables, start=1)]
                if not tables:
                    tables = [(Path(name).stem, pd.DataFrame({"Text": [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]}))]
            else:
                text = uploaded.getvalue().decode("utf-8", errors="replace")
                tables = [(Path(name).stem, pd.DataFrame({"Text": [line for line in text.splitlines() if line.strip()]}))]

            for table_name, dataframe in tables:
                sheet_name = excel_sheet_name(str(table_name), used_sheet_names)
                if dataframe.empty and len(dataframe.columns) == 0:
                    dataframe = pd.DataFrame({"Message": ["No structured rows were detected in this source."]})
                dataframe.fillna("").to_excel(writer, sheet_name=sheet_name, index=False)
                worksheet = writer.book[sheet_name]
                worksheet.freeze_panes = "A2"
                worksheet.auto_filter.ref = worksheet.dimensions
                for cell in worksheet[1]:
                    header_font = copy(cell.font)
                    header_font.bold = True
                    header_font.color = "FFFFFF"
                    cell.font = header_font
                    header_fill = copy(cell.fill)
                    header_fill.fill_type = "solid"
                    header_fill.fgColor = "17324D"
                    cell.fill = header_fill
                for column_cells in worksheet.columns:
                    column_letter = column_cells[0].column_letter
                    longest = max(len(str(cell.value or "")) for cell in column_cells[:100])
                    worksheet.column_dimensions[column_letter].width = min(max(longest + 2, 12), 55)
                if "Description" in dataframe.columns:
                    description_column = dataframe.columns.get_loc("Description") + 1
                    for cells in worksheet.iter_cols(min_col=description_column, max_col=description_column, min_row=2):
                        for cell in cells:
                            cell_alignment = copy(cell.alignment)
                            cell_alignment.wrap_text = True
                            cell.alignment = cell_alignment
                sheets_created.append(sheet_name)
    if not sheets_created:
        raise ValueError("No tables or text could be extracted from the uploaded files.")
    output.seek(0)
    return output.getvalue(), sheets_created


def main() -> None:
    initialize_database()
    st.set_page_config(page_title="Praj PDF Converter", page_icon="📄", layout="wide")
    st.markdown(
        """
        <style>
        @import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;700&family=Space+Grotesk:wght@500;700&display=swap');
        :root { --ink: #17324d; --mint: #bfe8d0; --paper: #f7f8f5; }
        html, body, [class*="css"] { font-family: 'DM Sans', sans-serif; }
        h1, h2, h3 { font-family: 'Space Grotesk', sans-serif !important; color: var(--ink); }
        .hero { background: linear-gradient(135deg, #17324d 0%, #235d68 62%, #6bb49f 100%); padding: 2.2rem 2.4rem; border-radius: 12px; color: white; margin-bottom: 1.4rem; }
        .hero h1 { color: white !important; font-size: 2.7rem; margin: 0; }
        .hero p { color: #e7f5ed; font-size: 1.05rem; margin: .5rem 0 0; max-width: 660px; }
        .metric { background: #ffffff; border: 1px solid #e4e7ec; border-left: 4px solid #6bb49f; padding: 1rem; min-height: 86px; }
        .metric strong { display:block; color: var(--ink); font-size: 1.6rem; font-family: 'Space Grotesk'; }
        .metric span { color: #667085; font-size: .85rem; }
        section[data-testid="stFileUploaderDropzone"] { background: #fbfdfb; border: 1px dashed #6bb49f; }
        </style>
        <div class="hero"><h1>Praj PDF Converter</h1><p>Drop photos, office files, spreadsheets, and PDFs into one polished document. OCR included for image-based paperwork.</p></div>
        """,
        unsafe_allow_html=True,
    )

    with st.sidebar:
        st.header("Conversion settings")
        conversion_mode = st.radio("Conversion mode", ["Files to PDF", "Files to Excel"], index=0)
        include_ocr = st.toggle("Run OCR on images", value=True)
        default_output = "praj_tables.xlsx" if conversion_mode == "Files to Excel" else "praj_pdf_bundle.pdf"
        output_name = st.text_input("Output filename", value=default_output)
        st.caption("OCR requires the Tesseract desktop engine. The PDF is still created if it is unavailable.")
        st.divider()
        st.subheader("Backend activity")
        logs = recent_logs()
        if logs:
            for created_at, status, logged_name, file_count, ocr_completed, error_message in logs:
                status_label = "Success" if status == "success" else "Failed"
                st.caption(f"{status_label} · {created_at.replace('T', ' ')}")
                st.write(f"{logged_name} · {file_count} file(s) · {ocr_completed} OCR")
                if error_message:
                    st.caption(error_message)
        else:
            st.caption("No conversions logged yet.")

    if conversion_mode == "Files to Excel":
        excel_files = st.file_uploader(
            "Upload files to extract into Excel",
            type=[extension for extensions in SUPPORTED_TYPES.values() for extension in extensions],
            accept_multiple_files=True,
            key="excel_source_files",
            help="Extract tables from PDFs and images, or copy data from Excel, CSV, DOCX, TXT, and Markdown files.",
        )
        if not excel_files:
            st.info("Upload a PDF, image, spreadsheet, or document to create an Excel workbook.")
            return
        st.subheader("Excel extraction queue")
        for position, file in enumerate(excel_files, start=1):
            st.write(f"{position:02d}  **{file.name}** · {(file_kind(file.name) or 'unsupported').title()} · {file.size / 1024:.1f} KB")
        if st.button("Extract tables to Excel", type="primary", use_container_width=True):
            with st.spinner("Detecting tables and creating your Excel workbook..."):
                try:
                    excel_bytes, sheet_names = build_excel(excel_files)
                    st.session_state["excel_bytes"] = excel_bytes
                    st.session_state["excel_sheet_names"] = sheet_names
                    log_conversion("success", output_name, excel_files, False)
                except Exception as error:
                    log_conversion("failed", output_name, excel_files, False, error_message=str(error))
                    st.error(f"Excel extraction failed: {error}")
        if st.session_state.get("excel_bytes"):
            st.success("Excel workbook is ready.")
            st.caption(f"Created {len(st.session_state.get('excel_sheet_names', []))} worksheet(s)")
            st.download_button(
                "Download Excel workbook",
                st.session_state["excel_bytes"],
                file_name=output_name if output_name.lower().endswith(".xlsx") else f"{output_name}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
        return

    files = st.file_uploader(
        "Upload your source files",
        type=[extension for extensions in SUPPORTED_TYPES.values() for extension in extensions],
        accept_multiple_files=True,
        help="Supported: PDF, PNG/JPG/WebP/BMP/TIFF, XLSX/XLS/CSV, DOCX, TXT, and Markdown.",
    )

    if not files:
        st.info("Start by uploading one or more files. They will be assembled in the order shown below.")
        return

    counts = {kind: sum(file_kind(file.name) == kind for file in files) for kind in SUPPORTED_TYPES}
    columns = st.columns(4)
    for column, (label, count) in zip(columns, [("Files", len(files)), ("Images", counts["image"]), ("Office", counts["document"] + counts["spreadsheet"]), ("PDFs", counts["pdf"])]):
        column.markdown(f'<div class="metric"><strong>{count}</strong><span>{label}</span></div>', unsafe_allow_html=True)

    st.subheader("Assembly queue")
    for position, file in enumerate(files, start=1):
        kind = file_kind(file.name) or "unsupported"
        st.write(f"{position:02d}  **{file.name}** · {kind.title()} · {file.size / 1024:.1f} KB")

    if st.button("Build PDF", type="primary", use_container_width=True):
        with st.spinner("Converting files and building your PDF..."):
            try:
                pdf_bytes, converted = build_pdf(files, include_ocr)
                st.session_state["pdf_bytes"] = pdf_bytes
                st.session_state["converted"] = converted
                log_conversion("success", output_name, files, include_ocr, converted)
            except Exception as error:
                log_conversion("failed", output_name, files, include_ocr, error_message=str(error))
                st.error(f"Conversion failed: {error}")

    if st.session_state.get("pdf_bytes"):
        st.success("Your PDF is ready.")
        converted = st.session_state.get("converted", [])
        ocr_count = sum(bool(item.ocr_text and "could not run" not in item.ocr_text) for item in converted)
        st.caption(f"Converted {len(converted)} file(s) · OCR completed for {ocr_count} image(s)")
        st.download_button("Download PDF", st.session_state["pdf_bytes"], file_name=output_name if output_name.lower().endswith(".pdf") else f"{output_name}.pdf", mime="application/pdf", use_container_width=True)
        ocr_items = [item for item in converted if item.kind == "image" and item.ocr_text]
        if ocr_items:
            with st.expander("View extracted OCR text"):
                for item in ocr_items:
                    st.markdown(f"**{item.name}**")
                    st.text_area(f"OCR result for {item.name}", item.ocr_text, height=150, key=f"ocr-{item.name}")
                combined_ocr = "\n\n".join(f"{item.name}\n{'-' * len(item.name)}\n{item.ocr_text}" for item in ocr_items)
                st.download_button("Download OCR text", combined_ocr, file_name="praj_pdf_ocr.txt", mime="text/plain")


if __name__ == "__main__":
    main()